import csv, io, os

import openpyxl
from flask import Blueprint, request, jsonify, session
from database import get_db
from routes.auth import login_required, manager_required

prod_bp = Blueprint('products', __name__, url_prefix='/api')


@prod_bp.route('/products')
@login_required
def list_products():
    q = request.args.get('q', '')
    page = int(request.args.get('page', 1))
    per_page = int(request.args.get('per_page', 50))
    supplier_id = request.args.get('supplier_id', type=int)
    offset = (page - 1) * per_page
    
    with get_db() as db:
        if q:
            params = [f'%{q}%', f'%{q}%', f'%{q}%', q, q, f'%{q}%']
            if supplier_id:
                params.append(supplier_id)
            sup_cond = ' AND p.supplier_id=?' if supplier_id else ''
            count_row = db.execute(
                'SELECT COUNT(DISTINCT p.id) as cnt FROM products p LEFT JOIN variants v ON v.product_id=p.id '
                'WHERE p.name LIKE ? OR p.sku LIKE ? OR v.sku LIKE ? OR v.barcode=? OR p.barcode=? OR p.category LIKE ?' + sup_cond,
                params
            ).fetchone()
            total = count_row['cnt']
            rows = db.execute(
                'SELECT DISTINCT p.id, p.name, p.category, p.base_price, p.cost_price, p.sku, p.barcode, p.has_variants, p.low_stock, p.created_at, p.commission_class, p.supplier_id, s.name as supplier_name FROM products p '
                'LEFT JOIN suppliers s ON s.id=p.supplier_id '
                'LEFT JOIN variants v ON v.product_id=p.id '
                'WHERE p.name LIKE ? OR p.sku LIKE ? OR v.sku LIKE ? OR v.barcode=? OR p.barcode=? OR p.category LIKE ?' + sup_cond +
                ' ORDER BY p.name LIMIT ? OFFSET ?',
                params + [per_page, offset]
            ).fetchall()
        else:
            if supplier_id:
                count_row = db.execute('SELECT COUNT(*) as cnt FROM products WHERE supplier_id=?', (supplier_id,)).fetchone()
                total = count_row['cnt']
                rows = db.execute(
                    'SELECT p.*, s.name as supplier_name FROM products p '
                    'LEFT JOIN suppliers s ON s.id=p.supplier_id '
                    'WHERE p.supplier_id=? ORDER BY p.name LIMIT ? OFFSET ?',
                    (supplier_id, per_page, offset)
                ).fetchall()
            else:
                count_row = db.execute('SELECT COUNT(*) as cnt FROM products').fetchone()
                total = count_row['cnt']
                rows = db.execute(
                    'SELECT p.*, s.name as supplier_name FROM products p '
                    'LEFT JOIN suppliers s ON s.id=p.supplier_id '
                    'ORDER BY p.name LIMIT ? OFFSET ?',
                    (per_page, offset)
                ).fetchall()
        
        products = []
        if rows:
            ids = [p['id'] for p in rows]
            vars_by_prod = {}
            for i in range(0, len(ids), 500):
                chunk = ids[i:i+500]
                ph = ','.join('?' * len(chunk))
                for v in db.execute(f'SELECT * FROM variants WHERE product_id IN ({ph}) ORDER BY product_id, size, color', chunk).fetchall():
                    vars_by_prod.setdefault(v['product_id'], []).append(dict(v))
            last_cost_by_prod = {}
            for i in range(0, len(ids), 500):
                chunk = ids[i:i+500]
                ph = ','.join('?' * len(chunk))
                for r in db.execute(
                    f'SELECT v.product_id, r.cost FROM restock_log r JOIN variants v ON v.id=r.variant_id '
                    f'WHERE v.product_id IN ({ph}) AND r.cost>0 ORDER BY r.id',
                    chunk
                ).fetchall():
                    last_cost_by_prod[r['product_id']] = round(r['cost'], 2)
            for p in rows:
                p = dict(p)
                variants = vars_by_prod.get(p['id'], [])
                p['variants'] = variants
                p['total_stock'] = sum(v['stock'] for v in variants)
                p['is_low'] = any(v['stock'] <= p['low_stock'] for v in variants) if variants else p['total_stock'] <= p['low_stock']
                p['last_purchased_cost'] = last_cost_by_prod.get(p['id'])
                products.append(p)
        
        return jsonify({
            'items': products,
            'total': total,
            'page': page,
            'per_page': per_page,
            'total_pages': (total + per_page - 1) // per_page
        })


@prod_bp.route('/products/search')
@login_required
def search_products():
    q = request.args.get('q', '')
    supplier_id = request.args.get('supplier_id', type=int)
    with get_db() as db:
        query = ('SELECT p.*, s.name as supplier_name, v.id as vid, v.size, v.color, v.sku as v_sku, '
                 'v.barcode as v_barcode, v.price as v_price, v.stock '
                 'FROM products p JOIN variants v ON v.product_id=p.id '
                 'LEFT JOIN suppliers s ON s.id=p.supplier_id '
                 'WHERE (p.name LIKE ? OR p.sku LIKE ? OR v.sku LIKE ? OR v.barcode=? OR p.barcode=?)')
        params = [f'%{q}%', f'%{q}%', f'%{q}%', q, q]
        if supplier_id:
            query += ' AND p.supplier_id=?'
            params.append(supplier_id)
        query += ' ORDER BY p.name, v.size, v.color'
        rows = db.execute(query, params).fetchall()
    results = []
    seen = set()
    for r in rows:
        key = (r['id'], r['vid'])
        if key in seen:
            continue
        seen.add(key)
        results.append({
            'pid': r['id'],
            'vid': r['vid'],
            'name': r['name'],
            'category': r['category'],
            'size': r['size'],
            'color': r['color'],
            'sku': r['v_sku'] or r['sku'],
            'barcode': r['v_barcode'] or r['barcode'],
            'price': r['v_price'] if r['v_price'] is not None else (r['base_price'] or 0),
            'base_price': r['base_price'],
            'stock': r['stock'],
            'low_stock': r['low_stock'],
            'supplier_id': r['supplier_id'],
            'supplier_name': r['supplier_name'],
        })
    return jsonify(results)


@prod_bp.route('/products', methods=['POST'])
@login_required
@manager_required
def add_product():
    d = request.get_json()
    with get_db() as db:
        try:
            supplier_id = d.get('supplier_id')
            if supplier_id:
                supplier_id = int(supplier_id)
            cur = db.execute(
                'INSERT INTO products (name, category, base_price, cost_price, sku, barcode, has_variants, low_stock, commission_class, supplier_id, brand, make, color) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)',
                (d['name'], d.get('category', ''), float(d.get('base_price', 0)), float(d.get('cost_price', 0)),
                 d['sku'], d.get('barcode'), int(d.get('has_variants', 0)), int(d.get('low_stock', 5)),
                 d.get('commission_class') or None, supplier_id,
                 d.get('brand', ''), d.get('make', ''), d.get('color', ''))
            )
            pid = cur.lastrowid
            if not d.get('has_variants'):
                stock = int(d.get('stock', 0))
                db.execute('INSERT INTO variants (product_id, sku, stock, size) VALUES (?,?,?,?)',
                           (pid, d['sku'], stock, d.get('size', '')))
                if stock > 0:
                    vid = db.execute('SELECT id FROM variants WHERE product_id=?', (pid,)).fetchone()['id']
                    db.execute(
                        'INSERT INTO restock_log (variant_id, old_stock, new_stock, qty_added, cost, note, staff_name) VALUES (?,?,?,?,?,?,?)',
                        (vid, 0, stock, stock, float(d.get('cost_price', 0)), 'Opening stock', session.get('name', ''))
                    )
            return jsonify({'ok': True, 'id': pid})
        except Exception as e:
            return jsonify({'error': str(e)}), 400


@prod_bp.route('/products/<int:pid>', methods=['PUT'])
@login_required
@manager_required
def update_product(pid):
    d = request.get_json()
    with get_db() as db:
        supplier_id = d.get('supplier_id')
        if supplier_id:
            supplier_id = int(supplier_id)
        db.execute(
            'UPDATE products SET name=?, category=?, base_price=?, cost_price=?, sku=?, barcode=?, low_stock=?, commission_class=?, supplier_id=?, brand=?, make=?, color=? WHERE id=?',
            (d['name'], d.get('category', ''), float(d['base_price']), float(d.get('cost_price', 0)),
             d['sku'], d.get('barcode'), int(d.get('low_stock', 5)), d.get('commission_class') or None, supplier_id,
             d.get('brand', ''), d.get('make', ''), d.get('color', ''), pid)
        )
        if not d.get('has_variants'):
            db.execute('UPDATE variants SET size=?, sku=? WHERE product_id=?',
                       (d.get('size', ''), d['sku'], pid))
        return jsonify({'ok': True})


@prod_bp.route('/products/<int:pid>', methods=['DELETE'])
@login_required
@manager_required
def delete_product(pid):
    """Delete a product if it has no invoice or stock-log references."""
    with get_db() as db:
        sale_count = db.execute(
            'SELECT COUNT(DISTINCT s.id) as cnt FROM sales s '
            'JOIN sale_items si ON si.sale_id = s.id WHERE si.product_id=?',
            (pid,)
        ).fetchone()['cnt']
        purchase_count = db.execute(
            'SELECT COUNT(DISTINCT pi.id) as cnt FROM purchase_invoices pi '
            'JOIN purchase_invoice_items pii ON pii.invoice_id = pi.id WHERE pii.product_id=?',
            (pid,)
        ).fetchone()['cnt']
        return_count = db.execute(
            'SELECT COUNT(DISTINCT pr.id) as cnt FROM purchase_returns pr '
            'JOIN purchase_return_items pri ON pri.return_id = pr.id WHERE pri.product_id=?',
            (pid,)
        ).fetchone()['cnt']
        restock_count = db.execute(
            'SELECT COUNT(*) as cnt FROM restock_log rl '
            'JOIN variants v ON v.id = rl.variant_id WHERE v.product_id=?',
            (pid,)
        ).fetchone()['cnt']

        refs = []
        if sale_count > 0:
            refs.append(f'{sale_count} sale(s)')
        if purchase_count > 0:
            refs.append(f'{purchase_count} purchase invoice(s)')
        if return_count > 0:
            refs.append(f'{return_count} purchase return(s)')
        if restock_count > 0:
            refs.append(f'{restock_count} stock log entry(ies)')
        if refs:
            return jsonify({'error': 'Cannot delete: product appears in ' + ', '.join(refs)}), 400

        try:
            db.execute('DELETE FROM products WHERE id=?', (pid,))
        except Exception as e:
            return jsonify({'error': 'Cannot delete: product is referenced by other records in the system.'}), 400
        return jsonify({'ok': True})


@prod_bp.route('/products/import-template')
@login_required
def import_template():
    """Return a sample CSV template for bulk product import."""
    header = 'name,sku,barcode,cost_price,base_price,category,commission_class,low_stock,stock,supplier_id'
    sample = 'Sample Product,SMP001,,150,250,Apparel,Standard,5,20,1'
    output = io.StringIO()
    output.write(header + '\n' + sample + '\n')
    return output.getvalue(), 200, {'Content-Type': 'text/csv', 'Content-Disposition': 'attachment; filename=product_import_template.csv'}


@prod_bp.route('/products/import', methods=['POST'])
@login_required
@manager_required
def import_products():
    """Bulk import products from CSV or Excel file. Skips duplicate SKUs, reports errors."""
    file = request.files.get('file')
    if not file:
        return jsonify({'error': 'No file uploaded'}), 400

    filename = file.filename.lower() if file.filename else ''
    rows = []

    if filename.endswith('.xlsx'):
        wb = openpyxl.load_workbook(file, read_only=True)
        ws = wb.active
        data_rows = list(ws.iter_rows(values_only=True))
        if not data_rows:
            return jsonify({'error': 'Empty file'}), 400
        headers = [str(h).strip().lower() if h else '' for h in data_rows[0]]
        for row in data_rows[1:]:
            d = {}
            for i, h in enumerate(headers):
                val = row[i] if i < len(row) else None
                if val is not None:
                    val = str(val).strip()
                d[h] = val
            rows.append(d)
    elif filename.endswith('.csv'):
        content = file.read().decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(content))
        if not reader.fieldnames:
            return jsonify({'error': 'Empty or invalid CSV'}), 400
        reader.fieldnames = [h.strip().lower() for h in reader.fieldnames]
        for row in reader:
            d = {}
            for k, v in row.items():
                d[k.strip().lower()] = v.strip() if v else ''
            rows.append(d)
    else:
        return jsonify({'error': 'Unsupported file format. Use .csv or .xlsx'}), 400

    expected = {'name', 'sku'}
    if not rows:
        return jsonify({'error': 'No data rows found'}), 400

    header_cols = set(rows[0].keys())
    missing_required = expected - header_cols
    if missing_required:
        return jsonify({'error': f'Missing required columns: {", ".join(sorted(missing_required))}'}), 400

    created = 0
    skipped = 0
    errors = []

    with get_db() as db:
        existing_skus = set(
            row['sku'] for row in db.execute('SELECT sku FROM products').fetchall()
        )
        existing_barcodes = set(
            row['barcode'] for row in db.execute('SELECT barcode FROM products WHERE barcode IS NOT NULL').fetchall()
        )

        db.execute('BEGIN IMMEDIATE')
        for idx, row in enumerate(rows, start=2):
            name = (row.get('name') or '').strip()
            sku = (row.get('sku') or '').strip()
            if not name:
                errors.append(f'Row {idx}: name is required')
                continue
            if not sku:
                errors.append(f'Row {idx}: sku is required')
                continue
            if sku in existing_skus:
                skipped += 1
                continue

            barcode = (row.get('barcode') or '').strip() or None
            if barcode and barcode in existing_barcodes:
                errors.append(f'Row {idx}: barcode "{barcode}" already exists')
                continue

            try:
                cost_price = float(row.get('cost_price', 0) or 0)
            except (ValueError, TypeError):
                errors.append(f'Row {idx}: invalid cost_price "{row.get("cost_price")}"')
                continue
            try:
                base_price = float(row.get('base_price', 0) or 0)
            except (ValueError, TypeError):
                errors.append(f'Row {idx}: invalid base_price "{row.get("base_price")}"')
                continue
            try:
                low_stock = int(row.get('low_stock', 5) or 5)
            except (ValueError, TypeError):
                errors.append(f'Row {idx}: invalid low_stock "{row.get("low_stock")}"')
                continue
            try:
                stock = int(row.get('stock', 0) or 0)
            except (ValueError, TypeError):
                errors.append(f'Row {idx}: invalid stock "{row.get("stock")}"')
                continue

            category = (row.get('category') or '').strip()
            commission_class = (row.get('commission_class') or '').strip() or None
            supplier_id = row.get('supplier_id', '').strip() or None
            if supplier_id:
                try:
                    supplier_id = int(supplier_id)
                except (ValueError, TypeError):
                    errors.append(f'Row {idx}: invalid supplier_id "{row.get("supplier_id")}"')
                    continue

            try:
                cur = db.execute(
                    'INSERT INTO products (name, category, base_price, cost_price, sku, barcode, low_stock, commission_class, supplier_id) VALUES (?,?,?,?,?,?,?,?,?)',
                    (name, category, base_price, cost_price, sku, barcode, low_stock, commission_class, supplier_id)
                )
                pid = cur.lastrowid
                cur2 = db.execute(
                    'INSERT INTO variants (product_id, sku, stock) VALUES (?,?,?)',
                    (pid, sku, stock)
                )
                if stock > 0:
                    db.execute(
                        'INSERT INTO restock_log (variant_id, old_stock, new_stock, qty_added, cost, note, staff_name) VALUES (?,?,?,?,?,?,?)',
                        (cur2.lastrowid, 0, stock, stock, cost_price, 'Opening stock', session.get('name', ''))
                    )
                existing_skus.add(sku)
                if barcode:
                    existing_barcodes.add(barcode)
                created += 1
            except Exception as e:
                errors.append(f'Row {idx}: {str(e)}')

    return jsonify({'ok': True, 'created': created, 'skipped': skipped, 'errors': errors})


@prod_bp.route('/variants', methods=['POST'])
@login_required
@manager_required
def add_variant():
    d = request.get_json()
    with get_db() as db:
        try:
            cur = db.execute(
                'INSERT INTO variants (product_id, size, color, sku, barcode, price, stock) VALUES (?,?,?,?,?,?,?)',
                (d['product_id'], d.get('size', ''), d.get('color', ''),
                 d['sku'], d.get('barcode'), float(d['price']) if d.get('price') else None,
                 int(d.get('stock', 0)))
            )
            pid = d['product_id']
            db.execute('UPDATE products SET has_variants=1 WHERE id=?', (pid,))
            return jsonify({'ok': True, 'id': cur.lastrowid})
        except Exception as e:
            return jsonify({'error': str(e)}), 400


@prod_bp.route('/variants/<int:vid>', methods=['PUT'])
@login_required
@manager_required
def update_variant(vid):
    d = request.get_json()
    with get_db() as db:
        db.execute(
            'UPDATE variants SET size=?, color=?, sku=?, barcode=?, price=?, stock=? WHERE id=?',
            (d.get('size', ''), d.get('color', ''), d['sku'], d.get('barcode'),
             float(d['price']) if d.get('price') else None, int(d.get('stock', 0)), vid)
        )
        return jsonify({'ok': True})


@prod_bp.route('/variants/<int:vid>/stock', methods=['PUT'])
@login_required
@manager_required
def update_stock(vid):
    d = request.get_json()
    with get_db() as db:
        old = db.execute('SELECT stock FROM variants WHERE id=?', (vid,)).fetchone()
        new_stock = int(d['stock'])
        qty = new_stock - old['stock']
        db.execute('UPDATE variants SET stock=? WHERE id=?', (new_stock, vid))
        if qty > 0:
            cost = float(d.get('cost', 0))
            db.execute(
                'INSERT INTO restock_log (variant_id, old_stock, new_stock, qty_added, cost, note, staff_name) VALUES (?,?,?,?,?,?,?)',
                (vid, old['stock'], new_stock, qty, cost, d.get('note', ''),
                 session.get('name', ''))
            )
            pid = db.execute('SELECT product_id FROM variants WHERE id=?', (vid,)).fetchone()
            if pid and cost > 0:
                total_row = db.execute(
                    'SELECT COALESCE(SUM(r.qty_added * r.cost),0) as total_cost, COALESCE(SUM(r.qty_added),0) as total_qty '
                    'FROM restock_log r JOIN variants v ON v.id=r.variant_id '
                    'WHERE v.product_id=? AND r.cost > 0',
                    (pid['product_id'],)
                ).fetchone()
                if total_row and total_row['total_qty'] > 0:
                    avg = round(total_row['total_cost'] / total_row['total_qty'], 2)
                    db.execute('UPDATE products SET cost_price=? WHERE id=?', (avg, pid['product_id']))
        return jsonify({'ok': True})


@prod_bp.route('/products/low-stock')
@login_required
def low_stock():
    with get_db() as db:
        products = db.execute('SELECT * FROM products ORDER BY name').fetchall()
        result = []
        for p in products:
            variants = db.execute('SELECT * FROM variants WHERE product_id=?', (p['id'],)).fetchall()
            low = [dict(v) for v in variants if v['stock'] <= p['low_stock']]
            if low:
                result.append({'product': dict(p), 'variants': low})
        return jsonify(result)


@prod_bp.route('/restock-log')
@login_required
def restock_log():
    with get_db() as db:
        rows = db.execute(
            'SELECT r.*, v.sku, v.size, v.color, p.name as product_name '
            'FROM restock_log r JOIN variants v ON v.id=r.variant_id '
            'JOIN products p ON p.id=v.product_id ORDER BY r.id DESC LIMIT 100'
        ).fetchall()
        return jsonify([dict(r) for r in rows])


@prod_bp.route('/barcodes')
@login_required
def list_barcodes():
    with get_db() as db:
        variants = db.execute(
            'SELECT v.id, v.sku, v.barcode, v.size, v.color, p.name as product_name, p.base_price, v.price '
            'FROM variants v JOIN products p ON p.id=v.product_id '
            'ORDER BY p.name'
        ).fetchall()
        return jsonify([dict(v) for v in variants])


@prod_bp.route('/products/bulk-delete', methods=['POST'])
@login_required
@manager_required
def bulk_delete_products():
    d = request.get_json() or {}
    ids = d.get('ids', [])
    if not ids:
        return jsonify({'error': 'No product IDs provided'}), 400
    deleted = 0
    errors = []
    with get_db() as db:
        for pid in ids:
            sale_count = db.execute(
                'SELECT COUNT(DISTINCT s.id) as cnt FROM sales s '
                'JOIN sale_items si ON si.sale_id = s.id WHERE si.product_id=?', (pid,)
            ).fetchone()['cnt']
            purchase_count = db.execute(
                'SELECT COUNT(DISTINCT pi.id) as cnt FROM purchase_invoices pi '
                'JOIN purchase_invoice_items pii ON pii.invoice_id = pi.id WHERE pii.product_id=?', (pid,)
            ).fetchone()['cnt']
            return_count = db.execute(
                'SELECT COUNT(DISTINCT pr.id) as cnt FROM purchase_returns pr '
                'JOIN purchase_return_items pri ON pri.return_id = pr.id WHERE pri.product_id=?', (pid,)
            ).fetchone()['cnt']
            restock_count = db.execute(
                'SELECT COUNT(*) as cnt FROM restock_log rl '
                'JOIN variants v ON v.id = rl.variant_id WHERE v.product_id=?', (pid,)
            ).fetchone()['cnt']
            if sale_count or purchase_count or return_count or restock_count:
                refs = []
                if sale_count: refs.append(f'{sale_count} sale(s)')
                if purchase_count: refs.append(f'{purchase_count} purchase invoice(s)')
                if return_count: refs.append(f'{return_count} return(s)')
                if restock_count: refs.append(f'{restock_count} stock log(s)')
                p = db.execute('SELECT name FROM products WHERE id=?', (pid,)).fetchone()
                name = p['name'] if p else f'ID {pid}'
                errors.append(f'{name}: referenced in {", ".join(refs)}')
                continue
            try:
                db.execute('DELETE FROM products WHERE id=?', (pid,))
                deleted += 1
            except Exception as e:
                errors.append(f'ID {pid}: {str(e)}')
    return jsonify({'ok': True, 'deleted': deleted, 'errors': errors})


@prod_bp.route('/products/bulk', methods=['PUT'])
@login_required
@manager_required
def bulk_update_products():
    d = request.get_json() or {}
    ids = d.get('ids', [])
    fields = d.get('fields', {})
    if not ids:
        return jsonify({'error': 'No product IDs provided'}), 400
    if not fields:
        return jsonify({'error': 'No fields to update'}), 400
    allowed = {'base_price', 'cost_price', 'category', 'commission_class', 'supplier_id', 'low_stock'}
    updates = []
    params = []
    for key, value in fields.items():
        if key not in allowed:
            continue
        if value == '' or value is None:
            if key in ('supplier_id', 'commission_class'):
                updates.append(f'{key}=NULL')
            else:
                updates.append(f'{key}=?')
                params.append(value)
        elif key in ('base_price', 'cost_price', 'low_stock'):
            updates.append(f'{key}=?')
            params.append(float(value) if key != 'low_stock' else int(value))
        elif key == 'supplier_id':
            updates.append(f'{key}=?')
            params.append(int(value) if value else None)
        else:
            updates.append(f'{key}=?')
            params.append(str(value))
    if not updates:
        return jsonify({'error': 'No valid fields to update'}), 400
    placeholders = ','.join('?' for _ in ids)
    sql = f'UPDATE products SET {",".join(updates)} WHERE id IN ({placeholders})'
    params.extend(ids)
    with get_db() as db:
        db.execute(sql, params)
    return jsonify({'ok': True, 'updated': len(ids)})


@prod_bp.route('/commission-classes')
@login_required
def list_commission_classes():
    with get_db() as db:
        rows = db.execute('SELECT * FROM commission_classes ORDER BY name').fetchall()
        return jsonify([dict(r) for r in rows])


@prod_bp.route('/commission-classes', methods=['POST'])
@login_required
@manager_required
def add_commission_class():
    d = request.get_json()
    name = d.get('name', '').strip()
    if not name:
        return jsonify({'error': 'Name required'}), 400
    percentage = float(d.get('percentage', 0))
    with get_db() as db:
        try:
            cur = db.execute('INSERT INTO commission_classes (name, percentage) VALUES (?,?)', (name, percentage))
            return jsonify({'ok': True, 'id': cur.lastrowid})
        except Exception as e:
            return jsonify({'error': str(e)}), 400


@prod_bp.route('/commission-classes/<int:ccid>', methods=['PUT'])
@login_required
@manager_required
def update_commission_class(ccid):
    d = request.get_json()
    name = d.get('name', '').strip()
    if not name:
        return jsonify({'error': 'Name required'}), 400
    percentage = float(d.get('percentage', 0))
    with get_db() as db:
        try:
            db.execute('UPDATE commission_classes SET name=?, percentage=? WHERE id=?', (name, percentage, ccid))
            return jsonify({'ok': True})
        except Exception as e:
            return jsonify({'error': str(e)}), 400


@prod_bp.route('/commission-classes/<int:ccid>', methods=['DELETE'])
@login_required
@manager_required
def delete_commission_class(ccid):
    with get_db() as db:
        db.execute('DELETE FROM commission_classes WHERE id=?', (ccid,))
        return jsonify({'ok': True})
